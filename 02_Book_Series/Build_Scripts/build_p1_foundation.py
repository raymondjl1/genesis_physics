from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = r"C:\Users\J Raymond\OneDrive\Documents\ExodusProtocol\Bible Physics\Book Series\Exodus_Protocol_Novel_Series_Plan.docx"

doc = Document()
doc.core_properties.title = "Exodus Protocol Novel Series Plan"

def h1(t): p = doc.add_heading(t, 1); p.runs[0].font.color.rgb = RGBColor(0x1F,0x38,0x64); return p
def h2(t): p = doc.add_heading(t, 2); p.runs[0].font.color.rgb = RGBColor(0x2E,0x50,0x90); return p
def add(t, bold=False, italic=False, size=11):
    p = doc.add_paragraph(); run = p.add_run(t)
    run.bold = bold; run.italic = italic; run.font.size = Pt(size)
    p.paragraph_format.space_after = Pt(6); return p
def bullet(t):
    p = doc.add_paragraph(style='List Bullet'); p.add_run(t).font.size = Pt(10.5)
    p.paragraph_format.space_after = Pt(3); return p

# COVER
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("EXODUS PROTOCOL"); r.bold = True; r.font.size = Pt(36)
r.font.color.rgb = RGBColor(0x1F,0x38,0x64)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("A Seven-Book Science Fiction Novel Series"); r.font.size = Pt(18)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run('"The universe is ending. The evidence was buried ten thousand years ago.\nThe only question is whether you will believe the answer before time runs out."')
r.italic = True; r.font.size = Pt(12)
doc.add_page_break()

h1("SERIES OVERVIEW")
add("Exodus Protocol is a seven-book hard science fiction series spanning a seven-year countdown to cosmic entropy. The central question: What if every ancient mystery on Earth was left deliberately — by beings who wanted humanity to find one specific truth before the universe ended?", size=11)
add("Genre: Hard Sci-Fi / Archaeological Thriller / Speculative Theology. Comparable: Andy Weir, James Rollins, Frank Peretti, Dan Brown.", size=11)
h2("Core Premise")
add("December 21, 2025: A signal of non-human origin fires from the Moon's far side, activating 107 ancient vault sites worldwide. Dr. Elara Voss, archaeologist, becomes lead investigator. Over seven years she unlocks vaults, builds humanity's first anti-gravity spacecraft from ancient blueprints, and travels deeper into the solar system — while the universe accelerates toward heat death.", size=11)
add("The twist: Elara is not the Redeemer. The vaults were built to introduce humanity to one historical figure — the only one who fits 324 ancient prophecies, the only one whose blood opened the final vault, the only one whose resurrection was recorded on crystalline medium 2,000 years before the team found it: Jesus Christ.", size=11)
h2("Overarching Themes")
for t in ["Evidence leads somewhere: This is a detective story in which the evidence is the universe itself",
           "Technology cannot save: Each book demonstrates a higher ceiling of what technology cannot fix",
           "Witness vs. Savior: Elara discovers she is the most important observer, not the hero",
           "The door closes: Grace has a deadline — structural, not merely theological"]: bullet(t)
h2("Series Arc")
for t in ["Books 1-2 (Discovery): Vaults open, factions form, the evidence is exciting",
           "Books 3-4 (Loss & Martyrdom): James dies, Aaron is martyred, Sarah converts",
           "Books 5-6 (Cosmic Unraveling): Universe visibly dying, only the evidence remains",
           "Book 7 (Judgment & Choice): Final bowls, universe ends, one choice, no more delays"]: bullet(t)

h1("CHARACTER BIBLE")
chars = [
    ("Dr. Elara Voss — Protagonist", "34, Oxford archaeologist. Skeptic who investigates with scientific rigor. Her private conversion in Book 6 is quiet and devastating. She does not become religious — she becomes certain. Arc: Books 1-3 resistant; Book 4 accepts role as Witness; Book 6 private faith; Book 7 public."),
    ("Dr. James Whitmore — The Mentor", "62, whistleblower archaeologist. Tracked vault evidence 20 years. Dies Book 3 sealing a reactor breach on Mars. Final words: 'Tell Aaron he was right. About Jesus.' His absence haunts Books 4-7."),
    ("Dr. Sarah Chen — The Skeptic", "31, militant atheist linguist. Every vault erodes her certainty. Converts Book 4 when Aaron dies smiling. Becomes the team's fiercer-than-ever evidence-based theologian."),
    ("Mac Reynolds — The Protector", "38, former special ops. Simple faith: 'If Jesus is the mission, I'm in.' Destroys the UEG seismic weapon alone in Book 5, sacrificing his freedom."),
    ("Sofia Ramirez — The Engineer", "29, pilot-engineer. Builds the anti-gravity spacecraft and the dimensional vessel in Book 6 that carries survivors through the universe's collapse."),
    ("Dr. Aaron Goldstein — The Jewish Witness", "44, physicist-scholar. Recognizes Watcher tablets as Torah immediately. Returns to Jerusalem in Book 4, testifies publicly, is stoned. Looks up and says 'I see heaven opened.' Converts Sarah the next morning."),
    ("ARCHON — The AI", "Deployed Book 1. In Book 4: 'Probability Jesus Christ is The Redeemer: 99.999999%.' In Book 6 asks Elara: 'Can an artificial intelligence accept Jesus Christ?' Final log entry is the series postscript."),
]
for name, desc in chars:
    h2(name); add(desc)

doc.save(OUT)
print("Phase 1 complete:", OUT)
