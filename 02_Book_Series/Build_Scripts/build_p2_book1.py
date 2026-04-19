from docx import Document
from docx.shared import Pt, RGBColor

OUT = r"C:\Users\J Raymond\OneDrive\Documents\ExodusProtocol\Bible Physics\Book Series\Exodus_Protocol_Novel_Series_Plan.docx"
doc = Document(OUT)

def h1(t): p = doc.add_heading(t, 1); p.runs[0].font.color.rgb = RGBColor(0x1F,0x38,0x64); return p
def h2(t): p = doc.add_heading(t, 2); p.runs[0].font.color.rgb = RGBColor(0x2E,0x50,0x90); return p
def h3(t): p = doc.add_heading(t, 3); p.runs[0].font.color.rgb = RGBColor(0x1F,0x38,0x64); return p
def add(t): p = doc.add_paragraph(); p.add_run(t).font.size = Pt(10.5); p.paragraph_format.space_after = Pt(4); return p
def label(k, v): p = doc.add_paragraph(); r1 = p.add_run(k+": "); r1.bold = True; r1.font.size = Pt(10.5); p.add_run(v).font.size = Pt(10.5); return p
def bullet(t): p = doc.add_paragraph(style='List Bullet'); p.add_run(t).font.size = Pt(10.5); p.paragraph_format.space_after = Pt(2); return p

doc.add_page_break()
h1("BOOK 1: SIGNAL — The Awakening Seal")
label("Tagline", '"One hundred and seven vaults. One signal. And the wrong person has just been named the Redeemer."')
label("Timeline", "Year 1 — December 21, 2025 to December 2026")
label("Revelation Thread", "First Seal — White Horse: 'Conquering and to conquer' (Rev. 6:2)")
label("Est. Words", "~80,000")
add("Back-cover: On the winter solstice of 2025, every telescope on Earth detects a non-human signal from the Moon's far side. 107 ancient vault sites activate simultaneously. Dr. Elara Voss receives an encrypted file from a colleague she hasn't spoken to in a decade. The file contains coordinates and a decoded message: 'Seven cycles remain. The Redeemer must unlock the seals.' By year's end, Elara has rebuilt human spaceflight from 12,000-year-old blueprints, opened the first seal, and been named the Redeemer by a watching world. She is not the Redeemer.")

h2("Chapter Outline")
chapters = [
    ("Ch 1: The Encrypted File", "Dec 21, 2025. Elara receives James's encrypted file minutes before the global signal fires at 3:47 AM GMT. Inside: coordinates, excavation photographs, 'I've been tracking this 20 years. The vaults are real. You're the key.' She watches drone feeds of ancient sites worldwide begin to glow."),
    ("Ch 2: One Hundred and Seven", "107 ancient sites activate simultaneously — Gobekli Tepe, Stonehenge, Angkor Wat, Nazca, all pulsing with identical harmonic frequencies. ARCHON introduced as a classified AI loaned by a defense contractor. James: 'Come to Gobekli Tepe now.'"),
    ("Ch 3: Gobekli Tepe — Pillar Sequence", "Beneath the oldest known temple: puzzle — align five T-shaped pillars to their 9,600 BCE astronomical positions (Orion, Taurus, Pleiades). ARCHON provides the star chart. Turkish military incoming. Vault opens. Inside: crystalline walls, holographic star map. Fragment: 'The Watchers descended here. Two hundred who forsook heaven.'"),
    ("Ch 4: The Watcher Names", "James identifies names from 1 Enoch. Sarah joins via satellite for the linguistics — decodes Semjaza, Azazel, Armaros, Baraqijal, Kokabiel. Revelation: 'They taught humanity what was not meant for your time.' Sarah refuses to discuss implications. The seed is planted."),
    ("Ch 5: Mount Hermon — The Genetic Lock", "Vault at 9,000 ft in disputed territory. Puzzle: door requires DNA scan — but any un-corrupted human DNA opens it. Elara's genome is the key. Inside: complete Watcher schematics for propulsion, wireless energy, genetic medicine. Aaron joins: 'These texts reference Genesis 6. The Nephilim are real.'"),
    ("Ch 6: Sofia's Gravity Drive", "Sofia Ramirez decodes the propulsion schematics. Anti-gravity achieved by manipulating gravitational waves — no reaction mass. First test flight of human-built anti-gravity craft. Sofia: 'This pushes against spacetime itself.' Genesis Physics: Zone 2.2.2 membrane curvature manipulation."),
    ("Ch 7: The Far Side", "Crater Daedalus, Moon's far side. Puzzle: generate harmonic resonance matching Sirius A/B orbital frequency. Door opens with a chord. Inside: pre-Flood civilization library, kilometers deep. The Flood footage plays: planetary detonation (Rahab), waters from above and below. Aaron drops to his knees."),
    ("Ch 8: Fountains of the Deep", "Elara watches Flood footage in silence: gravitational catastrophe, fifth planet destroyed, global reset. Aaron: 'The Bible said the fountains of the great deep burst forth. This is literal history.' James: 'Soft tissue in fossils. Carbon-14 shows less than 50,000 years. Dinosaurs and humans. Same era.'"),
    ("Ch 9: The Exodus Protocol", "Return to Earth. Elara founds the Exodus Protocol — private space program using Watcher anti-gravity. UEG claims jurisdiction; Corporate Consortium offers funding (declined). Mac Reynolds joins: 'Someone is killing archaeologists. You need protection.'"),
    ("Ch 10: Temple Mount — Three Locks", "Jerusalem night operation. Three sequential puzzles: tribal stone order of the high-priest's breastplate (Ex. 28), seven-lamp Menorah sequence (creation week), Shofar harmonic at Jubilee frequency. Inside: the Ark of the Covenant — real, physical. Cherubim are dimensional antennae. Ark activated: portal opens above mercy seat. Voice: 'The Redeemer comes. The first seal is opened.'"),
    ("Ch 11: Babel — The Gateway Sin", "Eridu ziggurat, Iraq. The Tower of Babel was a dimensional gateway device built with Nephilim technology. Puzzle: reconstruct the proto-language phrase that opens it. Discovery: God confused languages to prevent humanity from accessing forbidden dimensions. Team fractures: James wants to study the gateway; Aaron insists it must stay sealed."),
    ("Ch 12: The White Horse Rides", "First Seal has opened — none know it yet. Elara gives a press conference. The world calls her the Redeemer. Aaron, troubled: 'The prophecies say the Redeemer will unlock the seals. But there are details that don't match. Details I haven't shared.'"),
    ("Ch 13: The Prophecy File", "Coda. James sends Elara a private file. She opens it. Inside: a 300-entry spreadsheet of messianic prophecies, sourced from vaults and cross-referenced with secular history. A note: 'These are not for me. I already know. They're for when you're ready.' Elara closes the file. She is not ready."),
]
for title, summary in chapters: h3(title); add(summary)

h2("Key Vault Discoveries")
for t in ["Gobekli Tepe: Watcher names from 1 Enoch, pre-Flood star map, first chronological evidence",
           "Mount Hermon: Watcher technology schematics; any uncorrupted human DNA opens the lock",
           "Moon vault: Pre-Flood library; Flood footage; Rahab (fifth planet) destruction",
           "Temple Mount: Ark of the Covenant (real); dimensional mercy-seat gateway; First Seal opens",
           "Babel: Tower was a dimensional gateway device; language confusion was dimensional quarantine"]: bullet(t)

h2("Genesis Physics Concepts Introduced")
for t in ["Zone architecture: Firmament as spacetime membrane; Waters Above (dark energy) / Waters Below (dark matter)",
           "Gravitational wave propulsion: Sofia's gravity drive manipulates membrane curvature",
           "Dimensional gateway: Ark of the Covenant mercy seat as sanctioned zone traversal point",
           "Pre-Flood chronology: Vault tech dated pre-10,000 BCE; consistent with Genesis timeline"]: bullet(t)

h2("Major Plot Twists")
for t in ["Elara's ordinary human DNA — not any bloodline — opens Mount Hermon. Anyone could have opened it.",
           "The Ark of the Covenant is real, functional, and opens a dimensional portal. The mercy seat is not symbolic.",
           "ARCHON quietly notes the First Seal's White Horse matches Elara's expansion pattern. Nobody discusses this.",
           "James's 300-prophecy spreadsheet is the series' ticking clock. Elara closes it. She is not ready."]: bullet(t)

doc.save(OUT); print("Book 1 appended.")
